"""Tests for DocxReader: numbered extraction + numPr + fallback."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from sop_bpmn.reader import DocxReader


EXAMPLE = Path(__file__).resolve().parent.parent / "examples" / "input_sop.docx"
RICH = Path(__file__).resolve().parent.parent / "examples" / "input_sop_rich.docx"


def _inject_numpr(paragraph, ilvl: int = 0, num_id: int = 1) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    numPr.append(ilvl_el)
    numPr.append(num_id_el)
    pPr.append(numPr)


def test_extracts_numbered_steps_from_example():
    steps = DocxReader().extract_steps(str(EXAMPLE))
    assert len(steps) == 6
    assert steps[0].original_number == "1."
    assert steps[0].text.startswith("Receive customer support")
    assert all(s.level == 0 for s in steps)


def test_fallback_when_no_numbering(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Just some prose.")
    doc.add_paragraph("No leading numbers anywhere.")
    p = tmp_path / "unnumbered.docx"
    doc.save(str(p))

    steps = DocxReader().extract_steps(str(p))
    assert len(steps) == 2
    assert all(s.original_number == "" for s in steps)
    assert steps[0].text == "Just some prose."


def test_skips_empty_paragraphs(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("1. first")
    doc.add_paragraph("")
    doc.add_paragraph("2. second")
    p = tmp_path / "with_blanks.docx"
    doc.save(str(p))

    steps = DocxReader().extract_steps(str(p))
    assert len(steps) == 2
    assert [s.text for s in steps] == ["first", "second"]


def test_extracts_auto_numbered_word_list(tmp_path: Path):
    """Paragraphs with inline numPr (auto-numbering) should be extracted via the numPr path."""
    doc = Document()
    doc.add_paragraph("Heading without numbering")  # skipped — no numbering, no numPr
    for line in ["First step", "Second step", "Third step"]:
        p = doc.add_paragraph(line)
        _inject_numpr(p, ilvl=0, num_id=1)
    out = tmp_path / "auto.docx"
    doc.save(str(out))

    steps = DocxReader().extract_steps(str(out))
    assert len(steps) == 3
    assert [s.text for s in steps] == ["First step", "Second step", "Third step"]
    assert [s.original_number for s in steps] == ["1.", "2.", "3."]
    assert all(s.level == 0 for s in steps)


def test_extracts_mixed_manual_and_auto_numbered():
    """Rich fixture: Sec 1+2 use manual numbering (13 steps), Sec 3 uses numPr (3 steps)."""
    steps = DocxReader().extract_steps(str(RICH))
    assert len(steps) == 16
    # Section 1+2 still extracted via regex
    assert steps[0].text == "Receive customer support email"
    # Section 3 picked up via numPr
    sec3 = steps[-3:]
    assert [s.text for s in sec3] == [
        "Open the ticket dashboard",
        "Filter by priority",
        "Assign top item to on-call agent",
    ]
    assert [s.original_number for s in sec3] == ["1.", "2.", "3."]


def test_numpr_nested_levels(tmp_path: Path):
    """Nested numPr levels should produce correct level + restart deeper counters on shallower advance."""
    doc = Document()
    p1 = doc.add_paragraph("Top one"); _inject_numpr(p1, ilvl=0)
    p2 = doc.add_paragraph("Sub a"); _inject_numpr(p2, ilvl=1)
    p3 = doc.add_paragraph("Sub b"); _inject_numpr(p3, ilvl=1)
    p4 = doc.add_paragraph("Top two"); _inject_numpr(p4, ilvl=0)
    p5 = doc.add_paragraph("Sub a again"); _inject_numpr(p5, ilvl=1)
    out = tmp_path / "nested.docx"
    doc.save(str(out))

    steps = DocxReader().extract_steps(str(out))
    assert [s.level for s in steps] == [0, 1, 1, 0, 1]
    assert [s.original_number for s in steps] == ["1.", "1.", "2.", "2.", "1."]
