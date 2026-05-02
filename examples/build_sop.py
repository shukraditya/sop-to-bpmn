"""Build the rich synthetic SOP fixture used to exercise parser edge cases.

Produces examples/input_sop_rich.docx with three sections:
    Section 1 — happy path (manual numbering, single conditional + 2 branches)
    Section 2 — edge cases (false-positive `if`, multi-step branch via "Then ...")
    Section 3 — auto-numbered list (inline numPr; no leading N. in text)

Sections 1+2 use manual numbering so they're picked up by the existing regex.
Section 3 injects `<w:numPr>` directly so the reader's `numPr` path can detect
auto-numbering (Phase 3 target). The numId is synthetic — Word would need a
matching list definition to render the numbers, but the reader only checks
for numPr presence to decide which extraction path to use.
"""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent / "input_sop_rich.docx"


def _add_inline_numpr(paragraph, ilvl: int = 0, num_id: int = 1) -> None:
    """Inject <w:numPr> into the paragraph so reader detection picks it up."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    numPr.append(ilvl_el)
    numPr.append(num_id_el)
    pPr.append(numPr)


def build() -> Path:
    doc = Document()

    doc.add_heading("Customer Support SOP (rich fixture)", level=1)

    doc.add_heading("Section 1 — Triage (happy path)", level=2)
    happy = [
        "1. Receive customer support email",
        "2. Check if the issue is billing-related",
        "3. If yes, assign to Billing Queue",
        "4. If no, assign to General Support Queue",
        "5. Send acknowledgment email to customer",
        "6. Close the triage step",
    ]
    for line in happy:
        doc.add_paragraph(line)

    doc.add_heading("Section 2 — Edge cases", level=2)
    edge = [
        "1. Notify the manager if available",
        "2. Receive feedback",
        "3. Determine whether to refund the customer",
        "4. If yes, process the refund",
        "5. Then send a confirmation email",
        "6. If no, archive the request",
        "7. Close the case",
    ]
    for line in edge:
        doc.add_paragraph(line)

    doc.add_heading("Section 3 — Auto-numbered list", level=2)
    auto = [
        "Open the ticket dashboard",
        "Filter by priority",
        "Assign top item to on-call agent",
    ]
    for line in auto:
        p = doc.add_paragraph(line)
        _add_inline_numpr(p, ilvl=0, num_id=1)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
