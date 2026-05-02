"""Read .docx and emit a List[RawStep] for the parser."""

import re
from typing import List, Optional, Tuple

from docx import Document

from models import RawStep

# Manually-typed numbering patterns, in order of specificity.
# Each pattern: (regex, level). Group(1) = original_number, group(2) = body text.
NUMBER_PATTERNS: List[Tuple[re.Pattern, int]] = [
    (re.compile(r"^(\d+\.)\s+(.*)"), 0),         # "1. text"
    (re.compile(r"^(\(\d+\))\s+(.*)"), 1),       # "(1) text"
    (re.compile(r"^([a-z]\))\s+(.*)", re.I), 1), # "a) text"
]

NUMPR_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _parse_numbering(text: str) -> Optional[Tuple[int, str, str]]:
    for pattern, level in NUMBER_PATTERNS:
        m = pattern.match(text)
        if m:
            return level, m.group(1), m.group(2).strip()
    return None


def _has_numpr(paragraph) -> bool:
    return paragraph._p.find(f".//{NUMPR_NS}numPr") is not None


def _read_ilvl(paragraph) -> Optional[int]:
    """Return the numPr indent level (0 if numPr present without ilvl), or None if no numPr."""
    numpr_el = paragraph._p.find(f".//{NUMPR_NS}numPr")
    if numpr_el is None:
        return None
    ilvl_el = numpr_el.find(f"{NUMPR_NS}ilvl")
    if ilvl_el is None:
        return 0
    val = ilvl_el.get(f"{NUMPR_NS}val")
    try:
        return int(val) if val is not None else 0
    except ValueError:
        return 0


class DocxReader:
    """Extract numbered steps from a .docx file.

    Walks each paragraph once and emits a RawStep via whichever path matches:
        Path 1: leading number in `paragraph.text` (manual numbering).
        Path 2: `pPr/numPr/ilvl` (Word auto-numbering — number lives in formatting).
    Paragraphs matching neither (prose, headings, blanks) are skipped.

    If no paragraph matched either path, fall back to paragraph-per-step at level 0.
    """

    def extract_steps(self, path: str) -> List[RawStep]:
        doc = Document(path)
        steps: List[RawStep] = []
        # One counter per ilvl. When a shallower level advances, deeper counters reset.
        numpr_counters: List[int] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            parsed = _parse_numbering(text)
            if parsed:
                level, orig, body = parsed
                steps.append(RawStep(level=level, text=body, original_number=orig))
                continue

            ilvl = _read_ilvl(p)
            if ilvl is not None:
                while len(numpr_counters) <= ilvl:
                    numpr_counters.append(0)
                numpr_counters[ilvl] += 1
                for i in range(ilvl + 1, len(numpr_counters)):
                    numpr_counters[i] = 0
                orig = f"{numpr_counters[ilvl]}."
                steps.append(RawStep(level=ilvl, text=text, original_number=orig))

        if steps:
            return steps

        all_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return [RawStep(level=0, text=t, original_number="") for t in all_text]


def _inspect(path: str) -> None:
    """Print a day-1 diagnostic on the .docx — paragraph text + numPr presence."""
    doc = Document(path)
    print(f"=== Inspecting {path} ===")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    auto_count = 0
    for i, p in enumerate(doc.paragraphs):
        text = p.text
        numpr = _has_numpr(p)
        if numpr:
            auto_count += 1
        print(f"  [{i:2}] numPr={numpr} text={text!r}")
    if auto_count:
        print(
            f"\nNOTE: {auto_count} paragraph(s) have numPr (auto-numbered) — "
            "extracted via the numPr path with synthesized numbering."
        )


if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "examples/input_sop.docx"

    _inspect(path)

    print()
    reader = DocxReader()
    steps = reader.extract_steps(path)
    print(f"Extracted {len(steps)} step(s):")
    for i, s in enumerate(steps):
        print(f"  [{i:2}] level={s.level} num={s.original_number!r} text={s.text!r}")
